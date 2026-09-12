"""Generate docs/Task2_Summary.xlsx — business + developer summary for Task 2 (EDA).

PepsiCo-styled (reuses scripts/branding.py). Sheets: Overview, Key findings,
Figures, Modelling implications, Checklist.

Run with:  make task2-summary   (or)   python scripts/make_task2_summary.py
"""

from __future__ import annotations

from pathlib import Path

import branding as B
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "Task2_Summary.xlsx"

TITLE_FONT = Font(bold=True, size=14, color=B.PEPSI_BLUE)
HEADER_FONT = Font(bold=True, size=11, color=B.WHITE)
HEADER_FILL = PatternFill("solid", fgColor=B.PEPSI_BLUE)
TITLE_FILL = PatternFill("solid", fgColor=B.WHITE)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DONE = PatternFill("solid", fgColor="E2EFDA")


def _logo(height: int) -> XLImage | None:
    if not B.LOGO_SMALL.exists():
        return None
    img = XLImage(str(B.LOGO_SMALL))
    w, h = B.image_size(B.LOGO_SMALL)
    img.width = max(1, round(height * w / h))
    img.height = height
    return img


def _table(
    ws: Worksheet, title: str, headers: list[str], rows: list[list[str]], widths: list[int]
) -> None:
    span = max(len(headers), 2)
    for col in range(1, span + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=span)
    c = ws.cell(row=1, column=2, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 34
    logo = _logo(30)
    if logo is not None:
        ws.add_image(logo, "A1")
    for col, (head, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=2, column=col, value=head)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 22
    for r, row in enumerate(rows, start=3):
        for cc, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=cc, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        if "Status" in headers and str(row[headers.index("Status")]).startswith(("Done", "✅")):
            for k in range(1, len(headers) + 1):
                ws.cell(row=r, column=k).fill = DONE
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def build() -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "Overview"
    _table(
        ws,
        "Task 2 — Exploratory Data Analysis",
        ["Field", "Detail"],
        [
            ["Task", "Task 2 of 19 — EDA on the synthetic dataset (local, Mac)"],
            [
                "Objective",
                "Build intuition for the demand signal before modeling; validate the data-generating process and derive feature hypotheses.",
            ],
            ["Inputs", "data/synthetic/demand.parquet (~985K rows, 960 series)"],
            [
                "Outputs",
                "Market-wise docs/EDA.xlsx (per-market + cross-market sheets), 7 figures in docs/eda/, findings in docs/eda_findings.md, and this workbook.",
            ],
            [
                "Reproducible",
                "python -m cpg_forecast.eda (or make eda) — regenerates data if missing.",
            ],
            [
                "Note",
                "EDA caught two generator bugs (GT not dominant; Australia weather negative) which were fixed — that is the point of EDA.",
            ],
            [
                "DISCLAIMER",
                "Independent learning project; NOT affiliated with/endorsed by PepsiCo; synthetic data.",
            ],
        ],
        [20, 104],
    )

    ws = wb.create_sheet("Key findings")
    _table(
        ws,
        "Key Findings",
        ["Finding", "Value", "Implication"],
        [
            [
                "Intermittency",
                "~1.6% zero-demand days",
                "Some sparse SKUs; count/quantile losses needed",
            ],
            [
                "Volume by market",
                "India > Pakistan > UK > Australia",
                "Market scale differs; market embeddings",
            ],
            [
                "Traditional/General Trade share",
                "India 52%, Pakistan 65% (dominant)",
                "Channel matters hugely; GT is no-POS (secondary sales)",
            ],
            ["Weekly seasonality", "weekend ≈ 1.15x weekday", "Add weekday features"],
            ["Promotion uplift", "≈ 1.59x vs non-promo", "Promo flag + price are key covariates"],
            ["Holiday uplift", "≈ 1.49x vs normal", "Per-market holiday calendars matter"],
            [
                "Weather (beverages)",
                "temp↔demand ≈ +0.88 to +0.90 all markets",
                "Temperature drives beverage demand",
            ],
            ["Autocorrelation", "spikes at lags 7/14/21/28", "Strong weekly cycle; use lags of 7"],
            ["Hemisphere flip", "Australia inverted vs UK/IN/PK", "Seasonality is market-specific"],
        ],
        [30, 40, 46],
    )

    ws = wb.create_sheet("Figures")
    _table(
        ws,
        "Figures (docs/eda/)",
        ["File", "Shows"],
        [
            ["fig_units_by_market.png", "Total units by market"],
            ["fig_channel_share.png", "Channel share of volume per market (GT dominance in IN/PK)"],
            ["fig_weekly.png", "Weekly seasonality (mean units by weekday)"],
            ["fig_seasonality_hemisphere.png", "Monthly seasonality — Australia flipped vs others"],
            ["fig_acf.png", "Autocorrelation of daily demand (weekly spikes)"],
            ["fig_promo_holiday.png", "Promotion and holiday uplift"],
            ["fig_weather_beverage.png", "Beverage demand vs temperature (India)"],
        ],
        [34, 78],
    )

    ws = wb.create_sheet("Modelling implications")
    _table(
        ws,
        "Modelling Implications (feed Task 3+)",
        ["Area", "Implication"],
        [
            [
                "Features",
                "lags + rolling stats, weekday & month, promo & holiday flags, temperature, price",
            ],
            ["Categoricals", "market / channel / retailer / brand / SKU embeddings"],
            [
                "Loss",
                "count/quantile losses (Poisson / Tweedie / pinball) for intermittent SKUs, not plain MSE",
            ],
            [
                "Architecture",
                "per-market seasonality + channel differences -> embeddings and/or segmented models",
            ],
            ["Validation", "rolling-origin backtest per market (seasonality/scale differ)"],
        ],
        [22, 90],
    )

    ws = wb.create_sheet("Checklist")
    _table(
        ws,
        "Task 2 — Checklist",
        ["Deliverable", "Status"],
        [
            ["EDA script (reproducible, tested helpers)", "Done ✅"],
            ["Distributions / intermittency", "Done ✅"],
            ["Weekly + annual seasonality (hemisphere)", "Done ✅"],
            ["Autocorrelation (ACF)", "Done ✅"],
            ["Promo / holiday / weather effects", "Done ✅"],
            ["Channel & hierarchy analysis", "Done ✅"],
            ["Market-wise EDA workbook (docs/EDA.xlsx)", "Done ✅"],
            ["Figures (docs/eda/) + findings doc", "Done ✅"],
            ["Fixed 2 generator bugs found via EDA", "Done ✅"],
            ["Task2_Summary.xlsx (this file)", "Done ✅"],
        ],
        [52, 12],
    )

    return wb


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH.relative_to(OUT_PATH.parents[1])} with sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
