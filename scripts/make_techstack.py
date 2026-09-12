"""Generate docs/TechStack.xlsx — the full project technology stack.

PepsiCo-styled (reuses scripts/branding.py). Lists every technology across the
project: what is installed today vs planned per task, its purpose, and whether it
runs on the Mac (light) or cloud GPU (heavy).

Run with:  make techstack   (or)   python scripts/make_techstack.py
"""

from __future__ import annotations

from pathlib import Path

import branding as B
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "TechStack.xlsx"

TITLE_FONT = Font(bold=True, size=14, color=B.PEPSI_BLUE)
HEADER_FONT = Font(bold=True, size=11, color=B.WHITE)
HEADER_FILL = PatternFill("solid", fgColor=B.PEPSI_BLUE)
TITLE_FILL = PatternFill("solid", fgColor=B.WHITE)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREEN = PatternFill("solid", fgColor="E2EFDA")  # available now
AMBER = PatternFill("solid", fgColor="FBF3E2")  # planned
NOW = {"Installed", "In use", "Done"}


def _logo_image(height: int) -> XLImage | None:
    if not B.LOGO_SMALL.exists():
        return None
    img = XLImage(str(B.LOGO_SMALL))
    w, h = B.image_size(B.LOGO_SMALL)
    img.width = max(1, round(height * w / h))
    img.height = height
    return img


def _title(ws: Worksheet, title: str, span: int) -> None:
    span = max(span, 2)
    for col in range(1, span + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=span)
    c = ws.cell(row=1, column=2, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 34
    logo = _logo_image(30)
    if logo is not None:
        ws.add_image(logo, "A1")


def _headers(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    for col, (head, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=2, column=col, value=head)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 22


# ---- Stack content: (Layer, [(Technology, Purpose, Status, Task, Runs on)]) ----
STACK: list[tuple[str, list[tuple[str, str, str, str, str]]]] = [
    (
        "Language & runtime",
        [
            (
                "Python 3.12",
                "Primary language; ML wheels support (3.14 too new)",
                "Installed",
                "Task 0",
                "Both",
            ),
            ("venv", "Isolated per-project environment", "Installed", "Task 0", "Both"),
            ("pip + setuptools", "Packaging; editable src/ install", "Installed", "Task 0", "Both"),
            (
                "pyproject.toml",
                "Single source of truth for deps + tool config",
                "Installed",
                "Task 0",
                "Both",
            ),
            (
                "Makefile",
                "One-command workflows (setup/lint/test/summary)",
                "Installed",
                "Task 0",
                "Mac",
            ),
        ],
    ),
    (
        "Core data",
        [
            ("NumPy", "Numerical arrays", "Installed", "Task 0", "Both"),
            ("pandas", "DataFrames / ETL", "Installed", "Task 0", "Both"),
            (
                "PyArrow (Parquet)",
                "Columnar on-disk storage for the dataset",
                "Installed",
                "Task 1",
                "Both",
            ),
            ("Matplotlib", "Plotting / EDA charts", "Installed", "Task 2", "Mac"),
        ],
    ),
    (
        "Config & validation",
        [
            (
                "pydantic",
                "Schema + config validation (data contract)",
                "Installed",
                "Task 1/3",
                "Both",
            ),
            ("python-dotenv", "Load env vars / secrets from .env", "Installed", "Task 1", "Both"),
        ],
    ),
    (
        "Code quality & testing",
        [
            ("ruff", "Linting + import sorting", "Installed", "Task 0", "Mac"),
            ("black", "Deterministic code formatting", "Installed", "Task 0", "Mac"),
            ("pre-commit", "Git hooks; quality gate before commit", "Installed", "Task 0", "Mac"),
            ("pytest", "Automated tests", "Installed", "Task 0", "Mac"),
            ("ipykernel", "Jupyter kernel for EDA notebooks", "Installed", "Task 2", "Mac"),
        ],
    ),
    (
        "Docs & reporting",
        [
            ("openpyxl", "Generate the per-task Excel summaries", "Installed", "Task 0", "Mac"),
            ("Pillow", "Generate brand icons / images in workbooks", "Installed", "Task 0", "Mac"),
            (
                "prompt-engineering / golden-prompt-patterns (skills)",
                "Prompt & instruction quality",
                "In use",
                "Task 0/17",
                "Mac",
            ),
            (
                "superpowers (plugin)",
                "brainstorming / verification workflows",
                "In use",
                "Task 0",
                "Mac",
            ),
            (
                "keynote-in-html (skill)",
                "Final portfolio presentation deck",
                "Planned",
                "Task 18",
                "Mac",
            ),
        ],
    ),
    (
        "Data & versioning",
        [
            (
                "Synthetic data generator (NumPy/pandas)",
                "Primary high-realism dataset (SKU x store x day)",
                "Planned",
                "Task 1",
                "Mac",
            ),
            (
                "Kaggle API + M5 dataset",
                "Optional real-data benchmark (proxy)",
                "Planned",
                "Task 11",
                "Both",
            ),
            ("DVC", "Data version control + remote", "Planned", "Task 1", "Both"),
        ],
    ),
    (
        "Classical baselines",
        [
            ("statsmodels", "ETS / ARIMA baselines", "Planned", "Task 4", "Mac"),
            ("pmdarima", "auto-ARIMA", "Planned", "Task 4", "Mac"),
            ("LightGBM", "Gradient-boosting baseline (strong bar)", "Planned", "Task 4", "Mac"),
        ],
    ),
    (
        "Deep learning",
        [
            ("PyTorch", "Core deep-learning framework", "Planned", "Task 5", "Mac (MPS) / Cloud"),
            (
                "PyTorch Lightning",
                "Training-loop framework",
                "Planned",
                "Task 5",
                "Mac (MPS) / Cloud",
            ),
            (
                "pytorch-forecasting",
                "TFT / DeepAR / N-BEATS implementations",
                "Planned",
                "Task 9/10",
                "Mac (MPS) / Cloud",
            ),
            ("darts", "Time-series model library", "Planned", "Task 6-10", "Mac (MPS) / Cloud"),
        ],
    ),
    (
        "Experiment tracking & HPO",
        [
            (
                "MLflow (or Weights & Biases)",
                "Experiment tracking + model registry",
                "Planned",
                "Task 4/11",
                "Both",
            ),
            ("Optuna", "Hyperparameter optimization", "Planned", "Task 11", "Mac / Cloud"),
        ],
    ),
    (
        "Model export & optimization",
        [
            (
                "ONNX + ONNX Runtime",
                "Portable, fast CPU inference graph",
                "Planned",
                "Task 13",
                "Both",
            ),
            ("TorchScript", "Serialized model alternative", "Planned", "Task 13", "Both"),
        ],
    ),
    (
        "Serving",
        [
            ("FastAPI", "Inference REST API (/predict, /health)", "Planned", "Task 14", "Cloud"),
            ("Uvicorn", "ASGI server", "Planned", "Task 14", "Cloud"),
            ("Docker", "Containerize the service", "Planned", "Task 14", "Both"),
        ],
    ),
    (
        "Deployment & CI/CD",
        [
            ("GitHub Actions", "CI/CD: lint, test, build, deploy", "Planned", "Task 15", "Cloud"),
            (
                "Container registry (GHCR / Docker Hub)",
                "Host the built image",
                "Planned",
                "Task 15",
                "Cloud",
            ),
            (
                "Cloud Run / Render / Fly.io",
                "Live endpoint hosting (non-Apple)",
                "Planned",
                "Task 15",
                "Cloud",
            ),
        ],
    ),
    (
        "Monitoring",
        [
            ("Evidently", "Data / concept drift monitoring", "Planned", "Task 16", "Cloud"),
            ("alibi-detect", "Drift / outlier detection", "Planned", "Task 16", "Cloud"),
        ],
    ),
    (
        "App / UI",
        [
            (
                "HTML / CSS / JS (inline SVG)",
                "Demand IQ product mockup",
                "Done",
                "Task 0/17",
                "Browser",
            ),
            ("Streamlit", "Interactive demo app for stakeholders", "Planned", "Task 17", "Both"),
            (
                "Claude API (Anthropic)",
                "LLM insight / plain-English narrative layer",
                "Planned",
                "Task 17",
                "Cloud",
            ),
        ],
    ),
    (
        "Compute & platform",
        [
            (
                "Mac M3 Pro (Apple Silicon)",
                "Local dev, EDA, baselines, smoke tests",
                "In use",
                "Task 0",
                "Mac",
            ),
            (
                "Kaggle / Google Colab GPU",
                "Heavy model training (non-Apple)",
                "Planned",
                "Task 5+",
                "Cloud (GPU)",
            ),
        ],
    ),
    (
        "Version control & collaboration",
        [
            ("git", "Version control", "Installed", "Task 0", "Both"),
            (
                "GitHub (personal, non-Apple)",
                "Remote repo; push per task",
                "In use",
                "Task 0",
                "Cloud",
            ),
            ("GitHub CLI (gh)", "Repo creation / management", "In use", "Task 0", "Mac"),
        ],
    ),
]


# ---- Mac-local feasibility of the full skills matrix (synthetic data) ----
# (Area, Technologies, Verdict, How on the Mac, Task)
LOCAL_FEAS: list[tuple[str, str, str, str, str]] = [
    ("Programming", "Python", "Fully local", "Python 3.12 venv", "Task 0"),
    (
        "Additional languages",
        "R, Scala, MATLAB",
        "Optional/soft",
        "R via brew; Scala on the JVM (with Spark); MATLAB is licensed -> skip",
        "Optional",
    ),
    (
        "Python ML production",
        "scikit-learn, ML pipelines",
        "Fully local",
        "sklearn Pipelines on synthetic data",
        "Task 4-5",
    ),
    (
        "Data science",
        "ML, statistics, business modeling",
        "Fully local",
        "numpy / scipy / sklearn + notebooks",
        "All",
    ),
    (
        "Python engineering",
        "OOP, modular programming",
        "Fully local",
        "src/ package layout (in progress)",
        "Task 0+",
    ),
    ("Testing", "pytest / unit testing", "Fully local", "pytest wired via pre-commit", "Task 0"),
    ("Code quality", "PEP8, Black", "Fully local", "ruff + black + pre-commit", "Task 0"),
    (
        "Scientific Python",
        "pandas, NumPy, SciPy, scikit-learn, Matplotlib",
        "Fully local",
        "core deps already installed",
        "Task 1-4",
    ),
    (
        "Classical time series",
        "ARIMA, SARIMA, SARIMAX",
        "Fully local",
        "statsmodels, pmdarima",
        "Task 4",
    ),
    (
        "Statistical forecasting",
        "Exponential Smoothing, ETS, state-space",
        "Fully local",
        "statsmodels",
        "Task 4",
    ),
    ("Prophet", "Prophet", "Fully local", "pip install prophet (cmdstanpy) — CPU", "Task 4"),
    ("Multivariate time series", "VAR", "Fully local", "statsmodels VAR", "Task 4"),
    (
        "ML forecasting",
        "XGBoost, LightGBM",
        "Fully local",
        "CPU training on synthetic data",
        "Task 4",
    ),
    (
        "Deep learning",
        "LSTM, GRU (PyTorch)",
        "Fully local",
        "PyTorch MPS (Metal) — small synthetic data trains on M3 Pro",
        "Task 6",
    ),
    (
        "Transformer forecasting",
        "TFT / Transformer TS",
        "Local-capable",
        "PyTorch MPS; small models; cloud only if scaling up",
        "Task 9",
    ),
    (
        "Time-series features",
        "lags, rolling stats, calendar, exogenous",
        "Fully local",
        "pandas feature pipeline",
        "Task 3",
    ),
    (
        "Trend & seasonality",
        "decomposition, seasonal patterns",
        "Fully local",
        "statsmodels STL",
        "Task 2",
    ),
    (
        "Anomaly detection",
        "TS anomaly / outlier detection",
        "Fully local",
        "statsmodels / sklearn / alibi-detect (CPU)",
        "Task 2/16",
    ),
    (
        "Backtesting",
        "rolling / expanding window",
        "Fully local",
        "custom harness + sklearn TimeSeriesSplit",
        "Task 4",
    ),
    (
        "Forecast metrics",
        "MAE, RMSE, MAPE, sMAPE, WAPE, WRMSSE",
        "Fully local",
        "metrics module",
        "Task 4",
    ),
    (
        "Bayesian modeling",
        "Bayesian inference, probabilistic",
        "Fully local",
        "PyMC / NumPyro — CPU sampling on small data",
        "Task 10",
    ),
    (
        "Probabilistic forecasting",
        "prediction distributions / uncertainty",
        "Fully local",
        "quantile loss, DeepAR (small), PyMC",
        "Task 10",
    ),
    (
        "Big data",
        "Apache Spark",
        "Local-capable",
        "PySpark local[*] mode + Java (brew) — single machine",
        "Task 3/12",
    ),
    (
        "Relational databases",
        "SQL, RDBMS",
        "Fully local",
        "DuckDB + SQLite (serverless); Postgres via Docker",
        "Task 1/3",
    ),
    ("NoSQL", "NoSQL databases", "Local-capable", "MongoDB via Docker (optional)", "Optional"),
    (
        "Azure ML",
        "Azure Machine Learning",
        "Cloud-only",
        "Substitute locally with MLflow (tracking + registry)",
        "Task 11 (sub)",
    ),
    (
        "Azure cloud",
        "Azure services",
        "Cloud-only",
        "Emulate locally; real deploy -> non-Apple cloud (Render / Cloud Run)",
        "Task 15",
    ),
    ("REST APIs", "FastAPI, Flask", "Fully local", "uvicorn on localhost", "Task 14"),
    (
        "Microservices",
        "REST services, integration",
        "Fully local",
        "docker-compose multi-service",
        "Task 14/15",
    ),
    (
        "Frontend collaboration",
        "React, Angular",
        "Local-capable",
        "Node + Vite React demo (HTML mockup already built)",
        "Task 17",
    ),
    ("Containers", "Docker", "Fully local", "Docker Desktop (Apple Silicon)", "Task 14"),
    (
        "Orchestration",
        "Kubernetes",
        "Local-capable",
        "kind / minikube / k3d single-node on the Mac",
        "Task 15",
    ),
    ("Version control", "Git", "Fully local", "git + GitHub", "Task 0"),
    (
        "CI/CD",
        "CI/CD pipelines",
        "Local-capable",
        "pre-commit now; GitHub Actions emulated locally with `act`",
        "Task 15",
    ),
    (
        "DevOps / MLOps",
        "Docker, Kubernetes, CI/CD, cloud",
        "Local-capable",
        "docker + local k8s + MLflow + act",
        "Task 12-16",
    ),
    (
        "Observability",
        "logging, monitoring, alerting",
        "Fully local",
        "structured logging; Prometheus + Grafana via docker-compose; Evidently",
        "Task 16",
    ),
    (
        "Security",
        "secure production practices",
        "Fully local",
        "FastAPI JWT/OAuth2, .env secrets, RBAC",
        "Task 14/15",
    ),
    (
        "End-to-end data products",
        "Data -> ML -> API -> app -> monitoring",
        "Fully local",
        "the whole pipeline runs on the Mac",
        "All",
    ),
    (
        "Production deployment",
        "model serving, deploy, monitor",
        "Local-capable",
        "local docker/k8s emulation; real endpoint -> non-Apple cloud",
        "Task 15",
    ),
    (
        "Product thinking",
        "architecture + business integration",
        "Optional/soft",
        "docs, Architecture sheet, README",
        "All",
    ),
    (
        "Professional skills",
        "communication, service orientation",
        "Optional/soft",
        "demonstrated via docs / deck",
        "All",
    ),
    (
        "Teamwork",
        "collaboration + autonomous delivery",
        "Optional/soft",
        "git workflow, branches, PRs",
        "All",
    ),
    ("Communication", "fluent English", "Optional/soft", "docs, deck, code comments", "All"),
]


def build() -> Workbook:
    wb = Workbook()

    # 1) Overview -----------------------------------------------------------
    ws = wb.active
    ws.title = "Overview"
    _title(ws, "PepsiCo Demand Forecasting — Technology Stack", 2)
    _headers(ws, ["Field", "Detail"], [24, 100])
    total = sum(len(items) for _, items in STACK)
    installed = sum(1 for _, items in STACK for _t, _p, s, _tk, _r in items if s in NOW)
    rows = [
        ["Project", "End-to-end CPG demand forecasting with time-series deep learning"],
        [
            "Stack size",
            f"{total} technologies across {len(STACK)} layers ({installed} available now, {total - installed} planned)",
        ],
        [
            "Philosophy",
            "Because the data is SYNTHETIC (small), the whole stack runs LOCAL-FIRST on the Mac — including deep learning via PyTorch MPS (Metal). Cloud is optional, only for scale or the real public deployment.",
        ],
        [
            "Mac-local (Task 1)",
            "Almost everything runs completely on the Mac. Cloud-only = Azure ML / Azure Cloud (substitute with MLflow + non-Apple cloud for real deploy); MATLAB skipped (licensed). See the 'Mac-Local Feasibility' sheet.",
        ],
        [
            "Guardrails",
            "No Apple ecosystem; personal GitHub. Mac stays comfortable because synthetic data is small; a --sample flag + optional cloud cover any scale-up.",
        ],
        [
            "Status legend",
            "Installed / In use / Done = available now (green). Planned = introduced at a later task (amber).",
        ],
        [
            "See also",
            "Mac-Local Feasibility sheet (this file); Task0_Summary.xlsx (scope/brands/channels/architecture); README.md (roadmap).",
        ],
    ]
    for r, row in enumerate(rows, start=3):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    # 2) Tech Stack ---------------------------------------------------------
    ws = wb.create_sheet("Tech Stack")
    _title(ws, "Full Technology Stack — by layer", 6)
    _headers(
        ws,
        ["Layer", "Technology", "Purpose", "Status", "Introduced", "Runs on"],
        [24, 30, 46, 11, 12, 13],
    )
    r = 3
    for layer, items in STACK:
        for j, (tech, purpose, status, task, runs) in enumerate(items):
            vals = [layer if j == 0 else "", tech, purpose, status, task, runs]
            for c, val in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.alignment = WRAP
                cell.border = BORDER
            ws.cell(row=r, column=4).fill = GREEN if status in NOW else AMBER
            ws.cell(row=r, column=4).alignment = WRAP_CENTER
            ws.row_dimensions[r].height = 28
            r += 1
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    # 3) Mac-Local Feasibility (from the skills matrix) --------------------
    ws = wb.create_sheet("Mac-Local Feasibility")
    _title(ws, "Mac-Local Feasibility — full skills matrix on synthetic data", 5)
    _headers(
        ws,
        [
            "Area / Stack",
            "Technologies",
            "Mac-local?",
            "How on the Mac (local tool / substitute)",
            "Task",
        ],
        [22, 26, 15, 50, 12],
    )
    verdict_fill = {
        "Fully local": GREEN,
        "Local-capable": PatternFill("solid", fgColor="DDEBF7"),
        "Cloud-only": AMBER,
        "Optional/soft": PatternFill("solid", fgColor="EDEDED"),
    }
    r = 3
    for area, techs, verdict, how, task in LOCAL_FEAS:
        for c, val in enumerate([area, techs, verdict, how, task], start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.cell(row=r, column=3).fill = verdict_fill.get(verdict, AMBER)
        ws.cell(row=r, column=3).alignment = WRAP_CENTER
        ws.row_dimensions[r].height = 30
        r += 1
    ffull = sum(1 for _a, _t, v, _h, _tk in LOCAL_FEAS if v == "Fully local")
    note = ws.cell(
        row=r,
        column=1,
        value=f"Verdict: with synthetic data, {ffull}/{len(LOCAL_FEAS)} areas run FULLY on the Mac and most of the rest are local-capable (Spark local mode, single-node Kubernetes via kind/minikube, CI via `act`, DL via PyTorch MPS). Only Azure ML / Azure Cloud are cloud-only (substitute: MLflow locally + a non-Apple cloud for the one real-deployment task). MATLAB is skipped (licensed).",
    )
    note.alignment = WRAP
    note.font = Font(italic=True, color=B.PEPSI_BLUE)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 58
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    return wb


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH.relative_to(OUT_PATH.parents[1])} with sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
