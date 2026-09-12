"""Generate docs/Roadmap.xlsx — the full Task 0 -> 18 project roadmap.

PepsiCo-styled (reuses scripts/branding.py). Three sheets:
  - Overview       : phases, status counts, legend
  - Roadmap        : Task | Name | Objective | DL concept | Dev skill | Runs on | Status
  - Task Flows     : Task | Name | Key sub-steps (flow) | Deliverable

Run with:  make roadmap   (or)   python scripts/make_roadmap.py
"""

from __future__ import annotations

from pathlib import Path

import branding as B
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "Roadmap.xlsx"

TITLE_FONT = Font(bold=True, size=14, color=B.PEPSI_BLUE)
HEADER_FONT = Font(bold=True, size=11, color=B.WHITE)
HEADER_FILL = PatternFill("solid", fgColor=B.PEPSI_BLUE)
TITLE_FILL = PatternFill("solid", fgColor=B.WHITE)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "Done": PatternFill("solid", fgColor="E2EFDA"),
    "Next": PatternFill("solid", fgColor="DDEBF7"),
    "Planned": PatternFill("solid", fgColor="FBF3E2"),
}


def _logo(height: int) -> XLImage | None:
    if not B.LOGO_SMALL.exists():
        return None
    img = XLImage(str(B.LOGO_SMALL))
    w, h = B.image_size(B.LOGO_SMALL)
    img.width = max(1, round(height * w / h))
    img.height = height
    return img


def _title(ws: Worksheet, title: str, span: int) -> None:
    span = max(span, 2)
    for col in range(1, span + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=span)
    c = ws.cell(row=1, column=2, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 34
    logo = _logo(30)
    if logo is not None:
        ws.add_image(logo, "A1")


def _headers(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    for col, (head, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=2, column=col, value=head)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 22


# ---- Roadmap data: (num, name, objective, flow, dl_concept, dev_skill, deliverable, runs, status)
ROADMAP: list[tuple[str, str, str, str, str, str, str, str, str]] = [
    (
        "0",
        "Foundations & environment",
        "Reproducible, industry-standard project skeleton (no models yet).",
        "Repo + src/ layout; Python 3.12 venv; pyproject deps; ruff+black+pre-commit; pytest smoke; Makefile; git + push to personal GitHub.",
        "— (tooling)",
        "Project layout, env isolation, linting, testing, VCS",
        "Green scaffold + repo pushed",
        "Mac",
        "Done",
    ),
    (
        "1",
        "Problem framing & synthetic data (DVC)",
        "Turn business pain into an ML spec; build a high-realism, channel-aware synthetic dataset.",
        "Problem spec (28d horizon, WRMSSE); canonical schema geo->channel->retailer->store->SKUxday; per-market calendars; DGP generator (trend/seasonality/promos/weather/intermittency/cold-start/hierarchy/noise); write Parquet; DVC-track; tests.",
        "Forecast task formulation; DGP design",
        "Data modeling, DVC, pydantic schema, DuckDB/SQL",
        "Versioned synthetic dataset + generator + spec",
        "Mac",
        "Done",
    ),
    (
        "2",
        "EDA & time-series understanding",
        "Build intuition for the signal before modeling.",
        "Distributions, intermittency; STL decomposition; ACF/PACF; promo/holiday/weather effects; channel & hierarchy analysis; feature hypotheses.",
        "Seasonality, stationarity, exogenous drivers",
        "Reproducible notebooks, plotting",
        "EDA notebook + findings doc",
        "Mac",
        "Done",
    ),
    (
        "3",
        "Data pipeline & feature engineering",
        "Deterministic raw -> model-ready tensors, leakage-free.",
        "Long format; lags/rolling/calendar/price/promo/weather features; channel/retailer/brand embedding prep; time-based splits + rolling-origin CV; windowing; PyTorch Dataset/DataLoader; (Spark local for scale).",
        "Windowing, categorical embeddings, leakage-free splits",
        "ETL as code, config (pydantic/Hydra), unit tests",
        "Feature pipeline + tests",
        "Mac",
        "Next",
    ),
    (
        "4",
        "Local baselines & eval harness",
        "Per-market baselines = the bar to beat, plus the scoring machinery everything reuses.",
        "WRMSSE + MAE/RMSE/MAPE/sMAPE/WAPE; naive/seasonal-naive/ETS/ARIMA/SARIMAX/Prophet/VAR/LightGBM/XGBoost; rolling-origin backtest; MLflow tracking; per-market leaderboard.",
        "Why baselines matter; backtesting",
        "Metrics, backtesting, MLflow",
        "Leaderboard v1 + reusable eval module",
        "Mac",
        "Planned",
    ),
    (
        "5",
        "MLP forecaster (first neural net)",
        "Cross into DL with the simplest network.",
        "Feed-forward over flattened window + embeddings; PyTorch Lightning loop (loss/optimizer/early-stop/LR schedule); dropout/weight-decay; compare vs baselines.",
        "MLPs, backprop, regularization, LR scheduling",
        "Lightning module, checkpointing, sweeps",
        "MLP model + leaderboard update",
        "Mac (MPS)",
        "Planned",
    ),
    (
        "6",
        "RNN -> LSTM -> GRU",
        "Model temporal dependencies explicitly.",
        "Vanilla RNN (see vanishing gradients) -> LSTM -> GRU; packing, teacher forcing; lookback ablation; gradient clipping; seeds.",
        "Recurrence, gating, vanishing/exploding gradients",
        "MPS training discipline, reproducibility",
        "Recurrent models + comparison writeup",
        "Mac (MPS)",
        "Planned",
    ),
    (
        "7",
        "Seq2Seq + Attention",
        "Multi-step-ahead (28d) forecasting done right.",
        "Encoder-decoder; teacher forcing vs free-running; add attention + visualize weights; recursive vs direct multi-step.",
        "Seq2seq, attention, exposure bias",
        "Attention viz, training-strategy experiments",
        "Seq2seq+attention model + attention plots",
        "Mac (MPS)",
        "Planned",
    ),
    (
        "8",
        "Temporal Convolutional Networks (TCN)",
        "A convolutional alternative to recurrence.",
        "Dilated causal convolutions; receptive-field math; residual blocks; kernel/dilation ablation; compare to LSTM (speed/accuracy).",
        "Causal/dilated convolution, receptive field, residuals",
        "Architecture ablation, param/FLOP budgeting",
        "TCN model + tradeoff notes",
        "Mac (MPS)",
        "Planned",
    ),
    (
        "9",
        "Transformer / TFT (global model)",
        "State-of-the-art attention model; test global vs local per market.",
        "Positional encoding; self-attention; Temporal Fusion Transformer (static/known-future/observed inputs) with market/channel/retailer/brand/SKU embeddings; variable-selection & attention interpretability; per-market comparison vs baselines.",
        "Self-attention, multi-head, entity embeddings, global vs local",
        "pytorch-forecasting, market-level comparison",
        "TFT model + interpretability report",
        "Mac (MPS) / Cloud (opt)",
        "Planned",
    ),
    (
        "10",
        "Probabilistic & specialized models",
        "Forecast distributions and handle intermittent demand.",
        "DeepAR intervals; quantile/pinball loss; N-BEATS/N-HiTS; Croston/Tweedie for intermittency; PyMC/NumPyro Bayesian; interval calibration/coverage checks.",
        "Probabilistic forecasting, quantiles, intermittency",
        "Calibration, model-family comparison",
        "Probabilistic models + calibration report",
        "Mac (MPS)",
        "Planned",
    ),
    (
        "11",
        "Tracking, HPO & model selection",
        "Pick the champion scientifically; optional real-data reality check.",
        "Optuna sweeps on top architectures; MLflow compare (WRMSSE + intervals + latency); champion + challenger decision; OPTIONAL M5 (Kaggle) real-data benchmark.",
        "HPO, bias/variance, model selection",
        "Optuna, experiment hygiene, decision docs",
        "model_selection.md + final leaderboard",
        "Mac / Cloud (opt)",
        "Planned",
    ),
    (
        "12",
        "Training-as-code + hybrid refinement",
        "Train the champion reproducibly; test hybrid architectures.",
        "Training CLI (Hydra config); segmented globals (Beverages vs Snacks / by channel); geography fine-tuning; hierarchical reconciliation (MinT); mixed precision; checkpoint/resume; architecture chosen by backtest.",
        "Segmented globals, fine-tuning, reconciliation",
        "Training-as-code, architecture selection by backtest",
        "Production model artifact + train.py",
        "Mac (MPS) / Cloud (opt)",
        "Planned",
    ),
    (
        "13",
        "Export & optimization",
        "Make the model portable and fast for CPU serving.",
        "Export to TorchScript/ONNX + verify parity; dynamic/int8 quantization (latency vs accuracy); package preprocessing + model into one inference class.",
        "Graph export, quantization, train/serve skew",
        "ONNX Runtime, perf benchmarking, packaging",
        "CPU-fast model bundle",
        "Mac",
        "Planned",
    ),
    (
        "14",
        "Serving API (FastAPI + Docker)",
        "A real inference service.",
        "FastAPI /predict (28d forecast + intervals) & /health; pydantic schemas; validation + error handling; JWT/OAuth2 auth; Dockerize (slim CPU image); integration tests.",
        "Inference-time preprocessing, interval outputs",
        "API design, containerization, security, testing",
        "Working container serving predictions",
        "Mac (Docker)",
        "Planned",
    ),
    (
        "15",
        "Cloud deployment + CI/CD",
        "A live public endpoint with automated delivery.",
        "Push image to registry (GHCR); deploy to non-Apple cloud (Cloud Run / Render); GitHub Actions (lint/test/build/deploy, emulate locally with `act`); single-node Kubernetes (kind/minikube) demo; smoke-test the live URL.",
        "— (production readiness)",
        "Cloud deploy, CI/CD, Kubernetes, secrets",
        "Live endpoint URL + green CI pipeline",
        "Cloud + Mac (local k8s)",
        "Planned",
    ),
    (
        "16",
        "Monitoring, drift & retraining",
        "Keep the model honest in production.",
        "Log predictions vs actuals; rolling WRMSSE; data/concept drift (Evidently/alibi-detect); alert thresholds; scheduled retraining + registry promotion; Prometheus + Grafana dashboard.",
        "Distribution shift, model decay, retraining cadence",
        "Observability, scheduling, model registry",
        "Monitoring dashboard + retraining runbook",
        "Mac (docker-compose) / Cloud",
        "Planned",
    ),
    (
        "17",
        "Demo app + LLM insight layer",
        "A stakeholder-facing experience.",
        "Streamlit app (forecast chart + intervals, filters market/channel/brand, what-if price/promo); optional React build (HTML mockup already done); LLM narrative layer (Claude API) crafted with the prompt-engineering skill; deploy.",
        "Communicating uncertainty; LLM-in-the-loop",
        "App UX, prompt design/testing, integration",
        "Deployed Streamlit demo",
        "Mac + Cloud",
        "Planned",
    ),
    (
        "18",
        "Docs, presentation & portfolio",
        "Package the whole journey for reviewers/recruiters.",
        "Polish README (architecture diagram, results, live links); MODEL_EVOLUTION.md; Keynote-HTML deck; demo walkthrough; final leaderboard + honest limitations.",
        "Synthesis — where each model family wins/loses",
        "Technical writing, presentation, storytelling",
        "Portfolio-ready repo + shareable HTML deck",
        "Mac",
        "Planned",
    ),
]

PHASES = [
    ("Foundations", "Tasks 0-1", "Scaffold + problem framing + synthetic data"),
    ("Understand & prepare", "Tasks 2-3", "EDA + leakage-free feature pipeline"),
    ("Baselines", "Task 4", "Local baselines = the bar + eval harness"),
    (
        "Deep-learning ladder",
        "Tasks 5-10",
        "MLP -> RNN/LSTM/GRU -> Seq2Seq+Attn -> TCN -> TFT -> probabilistic",
    ),
    ("Selection & scale", "Tasks 11-12", "HPO + champion selection + hybrid refinement"),
    ("Productionize", "Tasks 13-16", "Export -> serve -> deploy -> monitor"),
    ("Deliver", "Tasks 17-18", "Demo app + LLM layer + portfolio deck"),
]


def build() -> Workbook:
    wb = Workbook()
    done = sum(1 for r in ROADMAP if r[8] == "Done")
    nxt = sum(1 for r in ROADMAP if r[8] == "Next")

    # 1) Overview -----------------------------------------------------------
    ws = wb.active
    ws.title = "Overview"
    _title(ws, "PepsiCo Demand Forecasting — Roadmap (Task 0 -> 18)", 3)
    _headers(ws, ["Phase", "Tasks", "Focus"], [24, 14, 78])
    r = 3
    for phase, tasks, focus in PHASES:
        for c, val in enumerate([phase, tasks, focus], start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1
    for label, val in [
        ("Total tasks", f"{len(ROADMAP)} (Task 0 -> 18)"),
        ("Progress", f"{done} done, {nxt} next, {len(ROADMAP) - done - nxt} planned"),
        (
            "Where it runs",
            "Local-first on the Mac (DL via PyTorch MPS); cloud optional for scale / real deploy.",
        ),
        (
            "Data",
            "Synthetic-primary (channel-aware); Kaggle M5 as an optional real-data benchmark (Task 11).",
        ),
        (
            "Architecture",
            "Evidence-driven: Local baseline -> Global model -> per-market comparison -> Hybrid refinement.",
        ),
        ("See also", "Task0_Summary.xlsx (scope/brands/channels/architecture) and TechStack.xlsx."),
    ]:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True, color=B.PEPSI_BLUE)
        ws.cell(row=r, column=1).alignment = WRAP
        m = ws.cell(row=r, column=2, value=val)
        m.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        r += 1
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    # 2) Roadmap ------------------------------------------------------------
    ws = wb.create_sheet("Roadmap")
    _title(ws, "Roadmap — Task 0 to 18", 7)
    _headers(
        ws,
        ["Task", "Name", "Objective", "Core DL concept", "Engineering skill", "Runs on", "Status"],
        [6, 30, 40, 30, 30, 22, 10],
    )
    r = 3
    for num, name, obj, _flow, dl, skill, _deliv, runs, status in ROADMAP:
        for c, val in enumerate([num, name, obj, dl, skill, runs, status], start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        sc = ws.cell(row=r, column=7)
        sc.fill = STATUS_FILL.get(status, STATUS_FILL["Planned"])
        sc.alignment = WRAP_CENTER
        ws.row_dimensions[r].height = 46
        r += 1
    ws.freeze_panes = "C3"
    ws.sheet_view.showGridLines = False

    # 3) Task Flows & Deliverables -----------------------------------------
    ws = wb.create_sheet("Task Flows")
    _title(ws, "Task Flows & Deliverables — Task 0 to 18", 4)
    _headers(ws, ["Task", "Name", "Key sub-steps (flow)", "Deliverable"], [6, 30, 76, 34])
    r = 3
    for num, name, _obj, flow, _dl, _skill, deliv, _runs, _status in ROADMAP:
        for c, val in enumerate([num, name, flow, deliv], start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[r].height = 60
        r += 1
    ws.freeze_panes = "C3"
    ws.sheet_view.showGridLines = False

    return wb


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH.relative_to(OUT_PATH.parents[1])} with sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
