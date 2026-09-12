"""Generate docs/Task1_Summary.xlsx — business + developer summary for Task 1.

PepsiCo-styled (reuses scripts/branding.py). Sheets: Overview, What was done,
Dataset schema, DGP realism, Checklist.

Run with:  make task1-summary   (or)   python scripts/make_task1_summary.py
"""

from __future__ import annotations

from pathlib import Path

import branding as B
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

OUT_PATH = Path(__file__).resolve().parents[1] / "docs" / "Task1_Summary.xlsx"

TITLE_FONT = Font(bold=True, size=14, color=B.PEPSI_BLUE)
HEADER_FONT = Font(bold=True, size=11, color=B.WHITE)
HEADER_FILL = PatternFill("solid", fgColor=B.PEPSI_BLUE)
TITLE_FILL = PatternFill("solid", fgColor=B.WHITE)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DONE = PatternFill("solid", fgColor="E2EFDA")


def _logo(height: int) -> XLImage | None:
    if not B.LOGO_SMALL.exists():
        return None
    img = XLImage(str(B.LOGO_SMALL))
    w, h = B.image_size(B.LOGO_SMALL)
    img.width = max(1, round(height * w / h))
    img.height = height
    return img


def _title(ws: Worksheet, title: str, span: int) -> None:
    span = max(span, 2)
    for col in range(1, span + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=span)
    c = ws.cell(row=1, column=2, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 34
    logo = _logo(30)
    if logo is not None:
        ws.add_image(logo, "A1")


def _table(
    ws: Worksheet, title: str, headers: list[str], rows: list[list[str]], widths: list[int]
) -> None:
    _title(ws, title, len(headers))
    for col, (head, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=2, column=col, value=head)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 22
    for r, row in enumerate(rows, start=3):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = WRAP
            cell.border = BORDER
        if "Status" in headers and str(row[headers.index("Status")]).startswith(("Done", "✅")):
            for cc in range(1, len(headers) + 1):
                ws.cell(row=r, column=cc).fill = DONE
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def build() -> Workbook:
    wb = Workbook()

    ws = wb.active
    ws.title = "Overview"
    _table(
        ws,
        "Task 1 — Problem Framing & Synthetic Data",
        ["Field", "Detail"],
        [
            ["Task", "Task 1 of 19 — problem framing + synthetic dataset (local, Mac)"],
            [
                "Objective",
                "Turn the business pain into a measurable ML spec and build a realistic, channel-aware synthetic dataset to model on.",
            ],
            [
                "Prediction task",
                "Forecast units per SKU x store x day, 28-day horizon, with prediction intervals.",
            ],
            [
                "Hierarchy",
                "geography -> channel -> retailer -> store -> SKU x day (channel & retailer = static store attributes/embeddings).",
            ],
            [
                "Primary metric",
                "WRMSSE + MAE/RMSE/MAPE/sMAPE/WAPE; rolling-origin backtest per market.",
            ],
            [
                "Dataset (full)",
                "~991K rows | 960 series | 4 markets | 10 channels | 48 stores | 40 SKUs | 2022-2024 | 1.4 MB Parquet | generated in ~3s on Mac.",
            ],
            [
                "Why synthetic",
                "Full control + known irreducible noise floor; no Kaggle dependency; matches PepsiCo brands/markets/channels. M5 kept as an optional real-data benchmark (Task 11).",
            ],
            ["Where it runs", "Local on the Mac (CPU); no cloud needed for data."],
            [
                "DISCLAIMER",
                "Independent learning project; NOT affiliated with/endorsed by PepsiCo; synthetic data based on the public M5 proxy framing.",
            ],
        ],
        [22, 104],
    )

    ws = wb.create_sheet("What was done")
    _table(
        ws,
        "Task 1 — What Was Built",
        ["Area", "What was done", "Where"],
        [
            [
                "Problem spec",
                "Target, horizon, metrics, scope boundaries documented.",
                "docs/problem_spec.md",
            ],
            [
                "Canonical schema",
                "15-column contract + pydantic config + validation.",
                "src/cpg_forecast/data/schema.py",
            ],
            [
                "Master data",
                "Markets, brands (food/bev), channels+retailers, packs, scales.",
                "src/cpg_forecast/data/reference.py",
            ],
            [
                "Calendars",
                "Per-market holidays/festivals with demand lift (Diwali, Eid, Christmas...).",
                "src/cpg_forecast/data/calendars.py",
            ],
            [
                "Generator (DGP)",
                "Channel-aware SKU x store x day generator (seeded, reproducible).",
                "src/cpg_forecast/data/generator.py",
            ],
            [
                "CLI",
                "python -m cpg_forecast.data.generator [--sample] --out ...",
                "generator.main()",
            ],
            [
                "Tests",
                "Schema, determinism, no-NaN/negative, channel-in-market, hierarchy, promo.",
                "tests/test_data_generator.py",
            ],
            [
                "Dataset",
                "Full + sample Parquet written to data/synthetic/ (git-ignored).",
                "data/synthetic/",
            ],
            [
                "Data versioning",
                "DVC initialised; dataset tracked via demand.parquet.dvc (pointer in git, data out).",
                "dvc add / .dvc/",
            ],
        ],
        [22, 62, 40],
    )

    ws = wb.create_sheet("Dataset schema")
    _table(
        ws,
        "Canonical Schema — SKU x store x day",
        ["Column", "Type", "Description"],
        [
            ["date", "date", "Daily timestamp"],
            ["market", "str", "Geography: UK / Australia / India / Pakistan"],
            [
                "channel",
                "str",
                "Route-to-market (General Trade, Modern Trade, E-com, Q-com, Wholesale...)",
            ],
            ["retailer", "str", "Banner (Kirana Store, DMart, Reliance, Tesco, Woolworths...)"],
            ["store_id", "str", "Individual outlet (static: one market/channel/retailer)"],
            ["category", "str", "Food or Beverage"],
            ["brand", "str", "e.g. Pepsi, Lay's, Kurkure"],
            ["sku_id", "str", "brand + pack"],
            ["pack_size", "str", "e.g. 330ml Can, 45g"],
            ["base_price", "float", "List price (local currency units)"],
            ["price", "float", "Actual price on the day (promo-adjusted; known-future covariate)"],
            ["promo_flag", "int", "1 if on promotion (known-future covariate)"],
            ["holiday", "str", "Holiday/festival name or empty (known-future covariate)"],
            ["temp_c", "float", "Daily temperature proxy (exogenous)"],
            ["units", "int", "TARGET — units sold that day"],
        ],
        [16, 10, 84],
    )

    ws = wb.create_sheet("DGP realism")
    _table(
        ws,
        "Hidden Data-Generating Process — what the model must learn",
        ["Signal injected", "What the model must learn"],
        [
            ["Trend", "Slow growth/decline per series"],
            ["Weekly seasonality", "Weekend uplift + day-of-week pattern"],
            [
                "Annual seasonality (hemisphere flip)",
                "Summer/winter cycles; Australia is inverted vs UK/IN/PK",
            ],
            [
                "Holidays / festivals",
                "Market-specific peaks (Diwali, Eid, Christmas, Australia Day...)",
            ],
            ["Promotions + price discounts", "Uplift during promos and price elasticity"],
            ["Weather", "Beverage demand rises with temperature"],
            [
                "Channel effects",
                "GT (many small, no-POS) vs MT vs wholesale pack/volume differences",
            ],
            ["Intermittency", "Sparse/zero-heavy demand for slow SKUs"],
            ["Cold-start listings", "New SKUs that start partway through history"],
            ["Stockout censoring", "Occasional zero days that are not true zero demand"],
            ["Irreducible noise", "Poisson count noise — the theoretical accuracy floor"],
        ],
        [34, 76],
    )

    ws = wb.create_sheet("Checklist")
    _table(
        ws,
        "Task 1 — Checklist",
        ["Deliverable", "Status"],
        [
            ["Problem spec (target/horizon/metrics/scope)", "Done ✅"],
            ["Canonical schema + pydantic config + validation", "Done ✅"],
            ["Per-market holiday calendars", "Done ✅"],
            ["Channel-aware synthetic generator (DGP)", "Done ✅"],
            ["CLI + --sample mode", "Done ✅"],
            ["Tests passing (determinism, schema, hierarchy, promo)", "Done ✅"],
            ["Full + sample Parquet generated locally", "Done ✅"],
            ["Task1_Summary.xlsx (this file)", "Done ✅"],
            ["DVC init + track dataset", "Done ✅"],
        ],
        [58, 12],
    )

    return wb


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = build()
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH.relative_to(OUT_PATH.parents[1])} with sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
